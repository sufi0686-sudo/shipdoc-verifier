"""Get plain text out of attachments: .txt, PDF, Word and scanned images (OCR).

Heavy libraries are imported lazily so the basic (plain-text) stage runs even
if the PDF/OCR packages are not installed.
"""
import io
import os
import re
from typing import Optional, Tuple

from .adapters import Attachment

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MIN_PDF_TEXT_CHARS = 40  # below this we assume the PDF is a scan and OCR it


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _ocr_image(img) -> str:
    import pytesseract
    return pytesseract.image_to_string(img)


def _pdf_text(data: bytes) -> Tuple[str, Optional[str]]:
    import pdfplumber
    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    text = "\n".join(parts)
    if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
        return text, None
    # Probably a scan: rasterise each page and OCR it.
    try:
        from pdf2image import convert_from_bytes
        pages = convert_from_bytes(data, dpi=300)
        return "\n".join(_ocr_image(p) for p in pages), "ocr_used"
    except Exception as exc:  # poppler/tesseract missing
        return text, f"PDF has no text layer and OCR failed: {exc}"


def _docx_text(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip()
                if t and (not cells or cells[-1] != t):  # merged cells repeat
                    cells.append(t)
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(att: Attachment) -> Tuple[str, Optional[str]]:
    """Return (text, note). note is None on a clean read, else a warning/error."""
    ext = os.path.splitext(att.filename.lower())[1]
    if att.data is None:
        return att.text or "", None
    try:
        if ext == ".pdf":
            return _pdf_text(att.data)
        if ext == ".docx":
            return _docx_text(att.data), None
        if ext == ".doc":
            return "", "Legacy .doc is not supported; please send .docx or PDF"
        if ext in IMAGE_EXT:
            from PIL import Image
            return _ocr_image(Image.open(io.BytesIO(att.data))), "ocr_used"
        return _decode(att.data), None
    except Exception as exc:
        return "", f"Could not read {att.filename}: {exc}"


def detect_doc_type(filename: str, text: str) -> Optional[str]:
    """Guess whether a document is a Shipping Instruction ("SI") or Bill of Lading ("BL")."""
    fn = filename.lower()
    head = text[:600].lower()
    si = bl = 0
    if re.search(r"(^|[^a-z])(si|shipping[_\s-]?instructions?|instructions?)([^a-z]|$)", fn):
        si += 2
    if re.search(r"(^|[^a-z])(bl|b-l|bol|hbl|mbl|obl|bill[_\s-]?of[_\s-]?lading|draft[_\s-]?bl)([^a-z]|$)", fn):
        bl += 2
    if re.search(r"shipping\s+instructions?", head):
        si += 3
    if re.search(r"bill\s+of\s+lading|\bb/l\s*(no|number)|\bbl\s*(no|number)", head):
        bl += 3
    if si == bl:
        return None
    return "SI" if si > bl else "BL"
