"""Generate demo data: data/sample_inbox.json plus PDF/Word documents.

Each email carries an "_expected" block so run_batch.py can report local
accuracy before you submit to the real scoreboard.

    python scripts/make_samples.py
"""
import json
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
SAMPLES = os.path.join(DATA, "samples")

# ---------------------------------------------------------------- documents
SI_1 = """SHIPPING INSTRUCTION
Booking No: MYPKG260114
Shipper: Sinar Palm Oil Sdn. Bhd.
Lot 12, Jalan Kilang 4, 40000 Shah Alam, Selangor
Consignee: Rotterdam Food Imports B.V.
Waalhaven 22, 3089 Rotterdam
Notify Party: Same as consignee
Port of Loading: Port Klang, Malaysia
Port of Discharge: Rotterdam, Netherlands
Container Count: 2 x 40HC
Gross Weight (kg): 42,500.00
Description: Refined palm olein in flexibags
"""
BL_1 = """BILL OF LADING (DRAFT)
B/L No: MAEU260114RTM

SHIPPER
SINAR PALM OIL SDN BHD
LOT 12, JALAN KILANG 4, 40000 SHAH ALAM

CONSIGNEE
ROTTERDAM FOOD IMPORTS BV
WAALHAVEN 22, ROTTERDAM

NOTIFY PARTY
SAME AS CONSIGNEE

PORT OF LOADING: PORT KELANG (MYPKG)
PORT OF DISCHARGE: NLRTM
NO. OF CONTAINERS: TWO (2)
GROSS WEIGHT: 42500 KGS
"""

SI_2 = """Shipping Instructions - Booking PGU-88213
1. Exporter: Tech Components Malaysia Sdn Bhd
2. Consignee: Shenzhen Electronics Trading Co., Ltd.
3. Notify: Pacific Freight Agency Ltd
4. POL: Pasir Gudang
5. POD: Shenzhen, China
6. Qty of containers: 3
7. G.W.: 36,000 kg
"""
BL_2 = """BILL OF LADING
BL Number: ONEY8821300
| Shipper | TECH COMPONENTS MALAYSIA SDN. BHD. |
| Consignee | SHENZHEN ELECTRONICS TRADING CO LTD |
| Notify Party | PACIFIC FREIGHT AGENCY LIMITED |
| Port of Loading | MYPGU |
| Port of Discharge | YANTIAN |
| Number of Containers | 4 x 40' HC |
| Gross Weight | 38,000.00 KGS |
"""

SI_3 = """SHIPPING INSTRUCTION
Shipper / Exporter: Evergreen Rubber Products Sdn Bhd
Consignee: Midwest Industrial Supply LLC
Notify Party: Chicago Customs Brokers Inc.
Port of Loading: Penang
Port of Discharge: Los Angeles, USA
Number of containers: 1 x 20GP
Total Gross Weight: 22,046 LBS
"""
BL_3 = """BILL OF LADING
Shipper
Evergreen Rubber Products Sdn. Bhd.
Consignee
Midwest Industrial Supply, LLC
Notify Party
Chicago Customs Brokers, Inc.
Port of Loading: George Town (MYPEN)
Port of Discharge: USLAX
Containers: One (1)
Gross Weight: 10,000 KGS
"""

SI_4 = """SHIPPING INSTRUCTION
Shipper: Borneo Timber Exports Sdn Bhd
Consignee: Global Trade Partners LLC
Notify Party: Global Trade Partners LLC
Port of Loading: Port Klang
Port of Discharge: Jebel Ali
Container Count: 5
Gross Weight (kg): 110,250
"""
BL_4 = """BILL OF LADING
Shipper: BORNEO TIMBER EXPORTS SDN BHD
Consignee: GLOBAL TRADING PARTNERS LLC
Notify Party: GLOBAL TRADE PARTNERS LLC
Port of Loading: PORT KLANG
Port of Discharge: AEJEA
Container Count: 5
Gross Weight (kg): 110,250
"""

INVOICE = """INVOICE No. INV-2026-0912
Bill to: Sinar Palm Oil Sdn Bhd
Ocean freight Port Klang - Rotterdam  USD 3,400.00
THC origin                            MYR 1,150.00
Amount due: USD 3,650.00   Due date: 30 Sep 2026
"""


def inbox():
    return [
        {
            "email_id": "EM-001",
            "from": "docs@sinarpalm.com.my",
            "subject": "Please compare SI against draft BL - MYPKG260114",
            "body": "Hi team,\n\nAttached are our SI and the carrier's draft BL. Kindly check for "
                    "any discrepancies before we approve the BL.\n\nThanks,\nAina",
            "attachments": [{"filename": "SI_MYPKG260114.txt", "content": SI_1},
                            {"filename": "Draft_BL_MAEU260114RTM.txt", "content": BL_1}],
            "_expected": {"category": "doc_comparison", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-002",
            "from": "export@techcomp.my",
            "subject": "BL draft check - PGU-88213",
            "body": "Dear CS,\nPlease verify the draft bill of lading against our shipping "
                    "instruction and revert before finalisation.",
            "attachments": [{"filename": "shipping_instruction.txt", "content": SI_2},
                            {"filename": "bl_draft.txt", "content": BL_2}],
            "_expected": {"category": "doc_comparison", "mismatch_found": True,
                          "mismatched_fields": ["container_count", "gross_weight_kg"]},
        },
        {
            "email_id": "EM-003",
            "from": "logistics@evergreenrubber.my",
            "subject": "Cross-check SI vs BL for USLAX shipment",
            "body": "Hello, can you cross-check both documents attached? Weight on our SI is in pounds.",
            "attachments": [{"filename": "SI-USLAX.txt", "content": SI_3},
                            {"filename": "HBL-USLAX.txt", "content": BL_3}],
            "_expected": {"category": "doc_comparison", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-004",
            "from": "ops@borneotimber.my",
            "subject": "Compare documents before BL release",
            "body": "Team, please compare the attached SI and BL.",
            "attachments": [{"filename": "SI_BTE.txt", "content": SI_4},
                            {"filename": "BL_BTE.txt", "content": BL_4}],
            "_expected": {"category": "doc_comparison", "mismatch_found": True,
                          "mismatched_fields": ["consignee"]},
        },
        {
            "email_id": "EM-005",
            "from": "booking@carrier-line.com",
            "subject": "SI cut-off reminder - Vessel MSC AURORA V.126",
            "body": "Dear Shipper,\nKindly submit your shipping instructions for booking "
                    "KUL223190 before the SI cut-off on Thursday 12:00.",
            "attachments": [],
            "_expected": {"category": "si_request", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-006",
            "from": "billing@forwarder.com",
            "subject": "Invoice INV-2026-0912 for ocean freight",
            "body": "Please find attached our invoice. Kindly arrange payment by the due date.",
            "attachments": [{"filename": "INV-2026-0912.txt", "content": INVOICE}],
            "_expected": {"category": "invoice", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-007",
            "from": "cs@carrier-line.com",
            "subject": "Vessel delay update: ETA Rotterdam revised",
            "body": "Please be advised the vessel departed Port Klang 18 hours late due to port "
                    "congestion. Revised ETA Rotterdam is 14 Oct.",
            "attachments": [],
            "_expected": {"category": "general_update", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-008",
            "from": "promo@win-big.biz",
            "subject": "Congratulations! You've won a prize",
            "body": "Click here to claim your free gift. Limited time offer. Unsubscribe anytime.",
            "attachments": [],
            "_expected": {"category": "spam", "mismatch_found": False, "mismatched_fields": []},
        },
        {
            "email_id": "EM-009",
            "from": "docs@sinarpalm.com.my",
            "subject": "Please compare with the draft BL",
            "body": "Hi, SI attached. Please compare against the draft BL from the carrier.",
            "attachments": [{"filename": "SI_MYPKG260201.txt", "content": SI_1}],
            "_expected": {"category": "doc_comparison", "mismatch_found": False,
                          "mismatched_fields": [], "needs_human_review": True},
        },
    ]


def make_docx(path, title, text):
    import docx
    d = docx.Document()
    d.add_heading(title, level=1)
    rows = [l for l in text.strip().split("\n")[1:] if ":" in l]
    table = d.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for line in rows:
        k, v = line.split(":", 1)
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k.strip(), v.strip()
    d.save(path)


def make_pdf(path, text):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Courier", size=11)
    for line in text.strip().split("\n"):
        pdf.cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(path)


if __name__ == "__main__":
    os.makedirs(SAMPLES, exist_ok=True)
    with open(os.path.join(DATA, "sample_inbox.json"), "w") as fh:
        json.dump(inbox(), fh, indent=2)
    for name, content in [("SI_sample.txt", SI_2), ("BL_sample.txt", BL_2)]:
        with open(os.path.join(SAMPLES, name), "w") as fh:
            fh.write(content)
    try:
        make_docx(os.path.join(SAMPLES, "SI_sample.docx"), "Shipping Instruction", SI_2)
        pdf_text = "\n".join(": ".join(c.strip() for c in l.strip("| ").split("|")) if "|" in l else l
                             for l in BL_2.split("\n"))
        make_pdf(os.path.join(SAMPLES, "BL_sample.pdf"), pdf_text)
    except ImportError as exc:
        print(f"Skipped PDF/DOCX samples ({exc}). pip install python-docx fpdf2")
    print("Sample data written to", DATA)
